"""Read the monthly payroll workbook into Payslip records.

This module only *reads*. It never derives a payroll figure -- Gross, Net, EPF,
ETF and APIT are taken exactly as the sheet states them. If the sheet is wrong,
the payslip is wrong in the same way, and `models.consistency_warnings` is what
surfaces that. Silently "correcting" a number would be far worse.

Three things about these workbooks drive the design:

**The sheet is 6.8 MB of mostly nothing.** The sample spans 9138 rows x 150
columns but holds only ten populated rows; the rest is empty styled cells. A
plain `load_workbook()` takes minutes. So: `read_only=True` to stream, and stop
early rather than walking to the end.

**Headers are dirty.** Real header text in the sample includes `"Basic Salary "`,
`"NIC No "`, `"Internship Allowance /        Other All."` and the misspelling
`"Net Earinings/Payable"`. Matching on exact strings would be a trap, so
matching is normalised and every field accepts several spellings.

**The header is not on row 1.** Rows 1-2 are a summary band, row 3 is the real
header, rows 4-9 are employees, row 10 is totals. The header row is found by
scoring, not assumed.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .models import Payslip

#: The only sheet this app reads. A workbook may carry bank-upload tabs, branch
#: codes or prior months; all are ignored.
TARGET_SHEET = "Payment Summary"

#: How far down to hunt for the header before giving up.
HEADER_SEARCH_ROWS = 40

#: Minimum distinct columns a row must match to be believed as the header.
MIN_HEADER_MATCHES = 6

#: Stop after this many consecutive empty rows, so we never walk 9000 blanks.
BLANK_ROW_LIMIT = 25

#: A cell equal to one of these ends the employee block.
TOTALS_MARKERS = frozenset({
    "total paid", "total", "totals", "grand total", "total for the month",
})

MONTHS = (
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
)

#: Accepted header spellings per Payslip field. The first entry is only
#: documentation -- all are compared normalised, so spacing and case are
#: already handled and these cover genuinely different wordings.
#:
#: `net_salary` carries the sheet's typo ("Earinings") *and* the correct
#: spelling, so the reader keeps working if finance ever fixes it.
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    # identity
    "employee_number": ("EMP/EPF No:", "EMP/EPF No", "EPF No", "Employee Number"),
    "employee_name": ("Name", "Employee Name"),
    "designation": ("Designation", "Position"),
    # earnings
    "basic_salary": ("Basic Salary",),
    "attendance_allowance": ("Attendance Allowance",),
    "travelling_allowance": ("Travelling Allowance", "Travel Allowance"),
    "other_allowances": (
        "Internship Allowance / Other All.", "Internship Allowance",
        "Other Allowances", "Other All.",
    ),
    "over_time": ("Over Time", "Overtime", "OT"),
    "gross_pay": ("Gross Earnings", "Gross Pay", "Gross"),
    # deductions
    "no_pay": ("Nopay Amount", "No Pay Amount", "No Pay"),
    "apit_tax": ("APIT Tax", "APIT"),
    "salary_advance": ("Salary Advance", "Salary Adjustment", "Advance"),
    "epf_employee": ("EPF(8%)", "EPF (8%)", "E.P.F (8%)"),
    "total_deduction": ("Total Deductions", "Total Deduction"),
    "net_salary": (
        "Net Earinings/Payable",   # sic -- the sheet's own misspelling
        "Net Earnings/Payable", "Net Salary", "Net Payable",
    ),
    # employer contribution
    "epf_company": ("EPF (12%)", "EPF(12%)"),
    "etf_company": ("ETF (3%)", "ETF(3%)"),
}


class WorkbookError(Exception):
    """Base for every problem this reader reports to the user."""


class SheetNotFoundError(WorkbookError):
    pass


class HeaderNotFoundError(WorkbookError):
    pass


class NoEmployeeRowsError(WorkbookError):
    pass


def _normalise(text: object) -> str:
    """Fold a header into a comparable key.

    Lowercase, strip every whitespace character (not just the ends), so
    `"EPF (12%)"`, `"EPF(12%)"` and `"epf  (12%) "` all collapse to
    `"epf(12%)"`. Whitespace is the only difference these headers reliably
    have, and it is invisible in Excel -- so it is exactly what breaks
    exact-match code.
    """
    if text is None:
        return ""
    return re.sub(r"\s+", "", str(text)).strip().lower()


#: normalised alias -> field name
_ALIAS_LOOKUP: dict[str, str] = {
    _normalise(alias): field_name
    for field_name, aliases in COLUMN_ALIASES.items()
    for alias in aliases
}


def parse_period_from_filename(path: Path | str) -> tuple[str, str]:
    """Best-effort "August", "2026" from a filename. Blank if unclear.

    The sheet itself carries no month or year, so the filename is the only
    hint available. It is a *default* the user can override, never a fact --
    which is why an unparseable name returns blanks instead of raising.
    """
    stem = Path(path).stem
    month = next((m for m in MONTHS if re.search(rf"\b{m}\b", stem, re.I)), "")
    year_match = re.search(r"\b(19|20)\d{2}\b", stem)
    return month, (year_match.group(0) if year_match else "")


def _select_sheet(workbook, sheet_name: str = TARGET_SHEET):
    """Return the target worksheet, or fail loudly naming what was found.

    No fallback to "the first sheet" and no content sniffing: reading the wrong
    tab would yield plausible-looking payslips with wrong numbers, which is the
    worst failure mode this app has.
    """
    wanted = _normalise(sheet_name)
    for name in workbook.sheetnames:
        if _normalise(name) == wanted:
            return workbook[name]
    found = ", ".join(f'"{n}"' for n in workbook.sheetnames) or "none"
    raise SheetNotFoundError(
        f'This workbook has no "{sheet_name}" sheet. Sheets found: {found}.'
    )


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[int, str]]:
    """Locate the header by scoring candidate rows, not by assuming row 1.

    Rows 1-2 of the sample are a summary band that repeats about ten of the
    same labels, so "first row that looks like a header" picks the wrong one.
    Scoring by distinct fields matched picks row 3, which carries ~25.

    Returns (1-based row number, {column index -> field name}).
    """
    best_row, best_map = 0, {}
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS], start=1):
        mapping: dict[int, str] = {}
        for col, value in enumerate(row):
            field_name = _ALIAS_LOOKUP.get(_normalise(value))
            # First column wins a duplicate label, e.g. the repeated "EPF (12%)".
            if field_name and field_name not in mapping.values():
                mapping[col] = field_name
        if len(mapping) > len(best_map):
            best_row, best_map = index, mapping

    if len(best_map) < MIN_HEADER_MATCHES:
        raise HeaderNotFoundError(
            f"Could not find the payroll header in the first "
            f"{HEADER_SEARCH_ROWS} rows of the sheet. Matched only "
            f"{len(best_map)} known columns; expected at least "
            f"{MIN_HEADER_MATCHES}."
        )
    return best_row, best_map


#: Markers folded through the same normaliser as the cells they are compared
#: against. Built here rather than written pre-normalised so TOTALS_MARKERS
#: stays readable -- and so the two sides can never drift apart again.
_TOTALS_KEYS: frozenset[str] = frozenset(_normalise(m) for m in TOTALS_MARKERS)


def _is_totals_row(row: tuple) -> bool:
    """True for the "Total Paid" summary row that closes the employee block.

    Both sides of this comparison must go through `_normalise`. Comparing a
    normalised cell against a raw marker silently never matches, which lets the
    totals row through as a seventh "employee" whose salary is the sum of the
    other six.
    """
    return any(_normalise(v) in _TOTALS_KEYS for v in row if v is not None)


@dataclass
class ReadResult:
    """What the reader found, including what it could not make sense of."""

    slips: list[Payslip]
    sheet_name: str = ""
    header_row: int = 0
    source_path: str = ""
    unmapped_headers: list[str] = field(default_factory=list)

    def __len__(self) -> int:
        return len(self.slips)


def read_payslips(
    path: Path | str,
    month: str | None = None,
    year: str | None = None,
    sheet_name: str = TARGET_SHEET,
) -> ReadResult:
    """Read `Payment Summary` into Payslip records.

    `month`/`year` override the values guessed from the filename; pass them
    from the UI once the user confirms the period.
    """
    path = Path(path)
    if not path.exists():
        raise WorkbookError(f"File not found: {path}")

    file_month, file_year = parse_period_from_filename(path)
    month = month if month is not None else file_month
    year = year if year is not None else file_year

    # read_only streams the sheet instead of building 1.3M cell objects;
    # data_only gives cached results, since most rows hold pasted values with
    # no formula at all and we must never recompute.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = _select_sheet(workbook, sheet_name)
        rows = _read_rows(worksheet)
        header_row, column_map = _find_header_row(rows)
        slips = _build_slips(rows, header_row, column_map, month, year, path)
        unmapped = _unmapped_headers(rows[header_row - 1], column_map)
    finally:
        workbook.close()

    if not slips:
        raise NoEmployeeRowsError(
            f'Found the "{worksheet.title}" sheet and its header on row '
            f"{header_row}, but no employee rows below it.\n\n"
            "If this workbook was generated by a script rather than saved from "
            "Excel, its formula results may not be stored in the file. Open it "
            "in Excel and save it once, then try again."
        )

    return ReadResult(
        slips=slips,
        sheet_name=worksheet.title,
        header_row=header_row,
        source_path=str(path),
        unmapped_headers=unmapped,
    )


def _read_rows(worksheet) -> list[tuple]:
    """Stream rows, stopping after a long run of blanks.

    Without the blank-run cut-off this walks all 9138 rows of the sample, the
    overwhelming majority of which are empty cells that merely carry styling.
    """
    rows: list[tuple] = []
    blank_run = 0
    for row in worksheet.iter_rows(values_only=True):
        if all(v is None or str(v).strip() == "" for v in row):
            blank_run += 1
            if blank_run >= BLANK_ROW_LIMIT and rows:
                break
            rows.append(row)
            continue
        blank_run = 0
        rows.append(row)
    return rows


def _build_slips(rows, header_row, column_map, month, year, path) -> list[Payslip]:
    """Turn the rows below the header into Payslip records."""
    slips: list[Payslip] = []
    blank_run = 0

    for offset, row in enumerate(rows[header_row:], start=header_row + 1):
        if _is_totals_row(row):
            break

        values = {
            field_name: row[col]
            for col, field_name in column_map.items()
            if col < len(row)
        }
        if not any(v is not None and str(v).strip() != "" for v in values.values()):
            blank_run += 1
            if blank_run >= 3:
                break
            continue
        blank_run = 0

        slips.append(Payslip(
            month=month, year=year,
            source_ref=f"{path.name}!row {offset}",
            **values,
        ))
    return slips


def _unmapped_headers(header: tuple, column_map: dict[int, str]) -> list[str]:
    """Header labels we recognised as text but do not use.

    Surfaced so a renamed column shows up as "we ignored this" rather than
    quietly becoming a blank field on someone's payslip.
    """
    return [
        str(value).strip()
        for col, value in enumerate(header)
        if col not in column_map and value is not None and str(value).strip()
    ]
