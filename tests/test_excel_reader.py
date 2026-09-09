"""Tests for the payroll workbook reader.

Most tests build a small synthetic workbook so they run anywhere. The tests
against the real file are skipped when it is absent, since payroll data is
deliberately not committed to git.
"""

from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook

from salary_app.excel_reader import (
    HeaderNotFoundError,
    NoEmployeeRowsError,
    SheetNotFoundError,
    TARGET_SHEET,
    parse_period_from_filename,
    read_payslips,
)

REAL_WORKBOOK = Path(__file__).resolve().parent.parent / "Final Salary - August 2026 - Copy.xlsx"
needs_real_file = pytest.mark.skipif(
    not REAL_WORKBOOK.exists(), reason="payroll workbook not present (gitignored)"
)

# Header spelled exactly as the real sheet has it, trailing spaces, odd runs of
# whitespace, misspelling and all. If the reader copes with this it copes with
# the real file.
DIRTY_HEADER = [
    "No", "EMP/EPF No:", "Name", "Designation",
    "Basic Salary ", "Attendance Allowance", "Travelling Allowance",
    "Internship Allowance /        Other All.", "Over Time", "Gross Earnings",
    "EPF (12%)", "EPF(8%)", "ETF (3%)", "APIT Tax", "Salary Advance",
    "Nopay Amount", "Total Deductions", "Net Earinings/Payable",
]

EMPLOYEE_ROW = [
    2, 70, "Nilupul Kodikara", "AI Engineer",
    80000, 35000, 35000, None, None, 150000,
    9600, 6400, 2400, 0, None,
    0, 6400, 143600,
]


def build_workbook(tmp_path, rows, sheet_name=TARGET_SHEET, extra_sheets=()):
    """Write a workbook whose first sheet holds `rows`."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    for name in extra_sheets:
        wb.create_sheet(name)
    path = tmp_path / "Final Salary - August 2026.xlsx"
    wb.save(path)
    return path


# --- sheet selection ------------------------------------------------------

def test_reads_only_the_payment_summary_sheet(tmp_path):
    """Other sheets in the workbook must be ignored entirely."""
    wb = Workbook()
    decoy = wb.active
    decoy.title = "Bank Codes"
    decoy.append(DIRTY_HEADER)
    decoy.append([9, 999, "WRONG PERSON", "Decoy", 1, 1, 1,
                  None, None, 1, 1, 1, 1, 1, None, 1, 1, 1])

    target = wb.create_sheet(TARGET_SHEET)
    target.append(DIRTY_HEADER)
    target.append(EMPLOYEE_ROW)

    path = tmp_path / "multi.xlsx"
    wb.save(path)

    result = read_payslips(path)
    assert result.sheet_name == TARGET_SHEET
    assert len(result) == 1
    assert result.slips[0].employee_name == "Nilupul Kodikara"


def test_missing_sheet_fails_loudly_and_names_what_it_found(tmp_path):
    """Never fall back to the first sheet: wrong numbers would look plausible."""
    path = build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW],
                          sheet_name="Sheet1", extra_sheets=("Branch Codes",))
    with pytest.raises(SheetNotFoundError) as exc:
        read_payslips(path)
    message = str(exc.value)
    assert "Payment Summary" in message
    assert "Sheet1" in message and "Branch Codes" in message


def test_sheet_name_match_tolerates_case_and_spacing(tmp_path):
    """Excel tab names routinely carry invisible trailing spaces."""
    path = build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW],
                          sheet_name=" payment  summary ")
    assert len(read_payslips(path)) == 1


# --- header detection -----------------------------------------------------

def test_header_found_below_a_summary_band(tmp_path):
    """Rows 1-2 repeat some of the same labels; scoring must prefer the real row."""
    summary_band = ["Basic Salary ", "Gross Earnings", "EPF (12%)", "EPF(8%)"]
    rows = [summary_band, [910000, 1600000, 109200, 72800],
            DIRTY_HEADER, EMPLOYEE_ROW]
    result = read_payslips(build_workbook(tmp_path, rows))
    assert result.header_row == 3
    assert len(result) == 1
    assert result.slips[0].basic_salary == Decimal("80000")


def test_dirty_header_spellings_are_matched(tmp_path):
    """Trailing spaces, internal runs of spaces, and the sheet's own typo."""
    result = read_payslips(build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW]))
    slip = result.slips[0]
    assert slip.basic_salary == Decimal("80000")        # "Basic Salary "
    assert slip.net_salary == Decimal("143600")         # "Net Earinings/Payable"
    assert slip.epf_employee == Decimal("6400")         # "EPF(8%)" no space
    assert slip.epf_company == Decimal("9600")          # "EPF (12%)" with space


def test_corrected_typo_still_maps(tmp_path):
    """If finance ever fixes "Earinings", the reader must keep working."""
    header = list(DIRTY_HEADER)
    header[header.index("Net Earinings/Payable")] = "Net Earnings/Payable"
    result = read_payslips(build_workbook(tmp_path, [header, EMPLOYEE_ROW]))
    assert result.slips[0].net_salary == Decimal("143600")


def test_unrecognisable_sheet_raises(tmp_path):
    rows = [["Alpha", "Beta", "Gamma"], [1, 2, 3]]
    with pytest.raises(HeaderNotFoundError):
        read_payslips(build_workbook(tmp_path, rows))


def test_unmapped_headers_are_reported(tmp_path):
    """A column we ignore should be visible, not silently dropped."""
    header = DIRTY_HEADER + ["Some New Column"]
    result = read_payslips(build_workbook(tmp_path, [header, EMPLOYEE_ROW]))
    assert "Some New Column" in result.unmapped_headers


# --- row selection --------------------------------------------------------

def test_totals_row_is_excluded(tmp_path):
    """Regression: "Total Paid" normalises to "totalpaid".

    Comparing a normalised cell against a raw marker silently never matched,
    which admitted the totals row as a seventh employee whose salary was the
    sum of the other six.
    """
    totals = [None, None, "Total Paid", None,
              160000, 70000, 70000, None, None, 300000,
              19200, 12800, 4800, 0, None, 0, 12800, 287200]
    second = list(EMPLOYEE_ROW)
    second[2] = "Second Person"

    result = read_payslips(build_workbook(
        tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW, second, totals]))

    assert len(result) == 2
    names = [s.employee_name for s in result.slips]
    assert "Total Paid" not in names


def test_rows_below_the_totals_row_are_ignored(tmp_path):
    stray = list(EMPLOYEE_ROW)
    stray[2] = "Should Not Appear"
    rows = [DIRTY_HEADER, EMPLOYEE_ROW,
            [None, None, "Total Paid"] + [None] * 15, stray]
    result = read_payslips(build_workbook(tmp_path, rows))
    assert len(result) == 1


def test_no_employee_rows_raises_with_a_useful_hint(tmp_path):
    with pytest.raises(NoEmployeeRowsError) as exc:
        read_payslips(build_workbook(tmp_path, [DIRTY_HEADER]))
    assert "save it once" in str(exc.value).lower()


def test_source_ref_points_back_at_the_row(tmp_path):
    """So a wrong figure can be traced to its cell."""
    result = read_payslips(build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW]))
    assert "row 2" in result.slips[0].source_ref


# --- period ---------------------------------------------------------------

@pytest.mark.parametrize("name,expected", [
    ("Final Salary - August 2026 - Copy.xlsx", ("August", "2026")),
    ("payroll june 2025.xlsx", ("June", "2025")),
    ("Final Salary - December 2026.xlsx", ("December", "2026")),
    ("payroll.xlsx", ("", "")),
])
def test_parse_period_from_filename(name, expected):
    assert parse_period_from_filename(name) == expected


def test_period_can_be_overridden(tmp_path):
    path = build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW])
    result = read_payslips(path, month="March", year="2027")
    assert result.slips[0].period == "March 2027"


def test_period_defaults_from_filename(tmp_path):
    path = build_workbook(tmp_path, [DIRTY_HEADER, EMPLOYEE_ROW])
    assert read_payslips(path).slips[0].period == "August 2026"


# --- against the real workbook -------------------------------------------

@needs_real_file
def test_real_workbook_yields_six_employees():
    result = read_payslips(REAL_WORKBOOK)
    assert result.sheet_name == "Payment Summary"
    assert result.header_row == 3
    assert len(result) == 6


@needs_real_file
def test_real_workbook_nets_sum_to_the_sheets_own_total():
    """Independent check that we read the right rows and the right column."""
    result = read_payslips(REAL_WORKBOOK)
    assert sum(s.net_salary for s in result.slips) == Decimal("1408100")


@needs_real_file
def test_real_workbook_row_six_matches_the_june_payslip_employee():
    """Same person as the reference PDF: basic 80,000 and EPF 8% of 6,400."""
    slip = read_payslips(REAL_WORKBOOK).slips[2]
    assert slip.basic_salary == Decimal("80000")
    assert slip.epf_employee == Decimal("6400")
    assert slip.epf_company == Decimal("9600")
    assert slip.etf_company == Decimal("2400")
    assert slip.net_salary == Decimal("143600")
    assert slip.period == "August 2026"


@needs_real_file
def test_real_workbook_parses_quickly():
    """6.8 MB of mostly-empty styled cells must not stall the UI."""
    import time
    start = time.perf_counter()
    read_payslips(REAL_WORKBOOK)
    assert time.perf_counter() - start < 2.0
